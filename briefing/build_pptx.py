#!/usr/bin/env python3
"""Rebuild the 2026-10-01 SDA filling-test deck. Numbers come from the Ishida log photo."""

from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = "/workspace/briefing/SDA-20261001-filling-test.pptx"
MEDIA = "/workspace/briefing/media"
FONT = "Microsoft JhengHei"
W = [203.4,200.9,201.1,200.3,199.8,199.4,201.0,201.5,200.8,201.1,201.8,200.4,201.2,200.8,200.7,200.7,201.0,201.0,200.9,201.4,200.6,200.9,201.0,201.1,201.1,200.6,201.6,201.4,201.6,202.0,201.6,202.0,201.8,201.6,201.4,201.0,200.9,200.7,201.2,200.8,201.0,200.8,200.8,200.6,200.8,200.6,200.8,200.7,200.6,200.8]

NAVY = RGBColor(0x1B, 0x2A, 0x3D)
INK = RGBColor(0x1C, 0x19, 0x15)
MUTED = RGBColor(0x5E, 0x58, 0x4F)
ISSUE = RGBColor(0x9A, 0x3B, 0x12)
CAUSE = RGBColor(0x1E, 0x4D, 0x78)
OK = RGBColor(0x0E, 0x6B, 0x45)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
PAPER = RGBColor(0xF7, 0xF4, 0xEE)
ISSUE_BG = RGBColor(0xFF, 0xF1, 0xE8)
CAUSE_BG = RGBColor(0xEE, 0xF4, 0xF8)
OK_BG = RGBColor(0xE7, 0xF5, 0xEE)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "2026/10/01 SDA 1L 充填測試 問題與根因"
prs.core_properties.subject = "第一輪 50 袋"


def set_run(run, text, size, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", font)


def add_text(slide, l, t, w, h, paragraphs, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    tf = box.text_frame
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for i, p in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.space_after = Pt(p.get("after", 4))
        para.space_before = Pt(p.get("before", 0))
        para.level = p.get("level", 0)
        set_run(para.add_run(), p["text"], p.get("size", 14), p.get("bold", False), p.get("color", INK))
    return box


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    # no shadow
    spPr = sh._element.spPr
    effect = spPr.find(qn("a:effectLst"))
    if effect is not None:
        spPr.remove(effect)
    return sh


def blank():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, PAPER)
    return s


def crop_fill(path, box_w, box_h, dest):
    im = Image.open(path).convert("RGB")
    target = box_w / box_h
    src = im.width / im.height
    if src > target:
        nw = int(im.height * target)
        x = (im.width - nw) // 2
        im = im.crop((x, 0, x + nw, im.height))
    else:
        nh = int(im.width / target)
        y = max(0, (im.height - nh) // 2)
        im = im.crop((0, y, im.width, y + nh))
    im.save(dest, quality=88)
    return dest


def footer(s, n, label):
    add_text(s, 0.4, 7.12, 9, 0.28, [{"text": label, "size": 10, "color": MUTED}])
    add_text(s, 11.4, 7.12, 1.5, 0.28, [{"text": f"{n:02d}", "size": 10, "color": MUTED, "align": PP_ALIGN.RIGHT}])


def kicker(s, text):
    add_text(s, 0.42, 0.22, 10, 0.28, [{"text": text, "size": 11, "bold": True, "color": MUTED}])


def title(s, text, top=0.46):
    add_text(s, 0.42, top, 12.4, 0.55, [{"text": text, "size": 26, "bold": True}])


# ---------- chart ----------
def make_chart(path):
    font = ImageFont.truetype("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 18)
    font_s = ImageFont.truetype("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 15)
    im = Image.new("RGB", (1800, 520), (247, 244, 238))
    d = ImageDraw.Draw(im)
    L, T, R, B = 70, 30, 20, 40
    pw, ph = 1800 - L - R, 520 - T - B
    ymin, ymax = 198.6, 204.0

    def X(i):
        return L + i * pw / 49

    def Y(v):
        return T + (ymax - v) / (ymax - ymin) * ph

    d.line([(L, Y(200)), (1800 - R, Y(200))], fill=(196, 85, 42), width=2)
    d.line([(L, Y(201)), (1800 - R, Y(201))], fill=(217, 208, 195), width=1)
    for g in (199, 200, 201, 202, 203, 204):
        d.text((4, Y(g) - 10), str(g), font=font_s, fill=(138, 129, 117))
    pts = [(X(i), Y(v)) for i, v in enumerate(W)]
    d.line(pts, fill=(30, 77, 120), width=3)
    for i, v in enumerate(W):
        r = 6 if v > 201.5 or v < 200 else 4
        col = (196, 85, 42) if v > 201.5 else (30, 77, 120) if v < 200 else (27, 42, 61)
        d.ellipse((X(i) - r, Y(v) - r, X(i) + r, Y(v) + r), fill=col)
    d.text((X(0) + 8, Y(203.4) - 26), "第1袋 203.4", font=font, fill=(154, 59, 18))
    for t in range(0, 50, 5):
        d.text((X(t) - 8, 490), str(t + 1), font=font_s, fill=(138, 129, 117))
    im.save(path, quality=90)


make_chart("/tmp/sda-chart.png")
crop_fill(f"{MEDIA}/line.jpg", 6.183, 7.5, "/tmp/ppt-line.jpg")
crop_fill(f"{MEDIA}/hobo.jpg", 3.3, 5.2, "/tmp/ppt-hobo.jpg")
crop_fill(f"{MEDIA}/lumps.jpg", 4.0, 2.5, "/tmp/ppt-lumps-wide.jpg")
crop_fill(f"{MEDIA}/grate-close.jpg", 4.0, 2.5, "/tmp/ppt-grate.jpg")
crop_fill(f"{MEDIA}/chute.jpg", 4.1, 2.5, "/tmp/ppt-chute.jpg")
crop_fill(f"{MEDIA}/lumps.jpg", 3.15, 5.6, "/tmp/ppt-lumps-tall.jpg")
crop_fill(f"{MEDIA}/chute.jpg", 3.5, 4.55, "/tmp/ppt-chute-tall.jpg")
crop_fill(f"{MEDIA}/grate-close.jpg", 3.3, 4.55, "/tmp/ppt-grate-tall.jpg")
crop_fill(f"{MEDIA}/lumps-wide.jpg", 3.3, 5.6, "/tmp/ppt-lumps-wide2.jpg")

# ---------- 1 cover ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 7.15, 7.5, NAVY)
rect(s, 7.15, 0, 6.183, 7.5, NAVY)
s.shapes.add_picture("/tmp/ppt-line.jpg", Inches(7.15), Inches(0), Inches(6.183), Inches(7.5))
add_text(s, 0.5, 0.7, 6.3, 0.3, [{"text": "FILLING TEST  ·  ISSUE / ROOT CAUSE", "size": 12, "bold": True, "color": RGBColor(0xD7, 0xB4, 0x8A)}])
add_text(s, 0.5, 1.5, 6.3, 2.1, [
    {"text": "SDA 1 L 充填測試", "size": 36, "bold": True, "color": WHITE, "after": 0},
    {"text": "第一輪 50 袋", "size": 36, "bold": True, "color": WHITE},
])
add_text(s, 0.5, 3.8, 6.2, 1.5, [{"text": "2026/10/01。四個現象是同一件事：有結塊的粉沒有乾淨離開秤斗，所以閘門會堵、壁上會沾，Ishida 秤到的重量沒有全部進袋。", "size": 16, "color": RGBColor(0xE6, 0xDF, 0xD2)}])
add_text(s, 0.5, 6.2, 2, 0.7, [{"text": "名義重量", "size": 12, "color": RGBColor(0xC9, 0xC1, 0xB4), "after": 0}, {"text": "200 g / 袋", "size": 16, "bold": True, "color": WHITE}])
add_text(s, 2.6, 6.2, 2, 0.7, [{"text": "現場時段", "size": 12, "color": RGBColor(0xC9, 0xC1, 0xB4), "after": 0}, {"text": "09:28 – 11:57", "size": 16, "bold": True, "color": WHITE}])
add_text(s, 4.7, 6.2, 2.2, 0.7, [{"text": "已排除", "size": 12, "color": RGBColor(0xC9, 0xC1, 0xB4), "after": 0}, {"text": "負荷元、室溫室濕", "size": 16, "bold": True, "color": WHITE}])

# ---------- 2 ruled out ----------
s = blank(); kicker(s, "先排除什麼"); title(s, "測試前，秤與環境都是合格的")
add_text(s, 0.42, 1.1, 12, 0.35, [{"text": "這兩個結果把根因從儀器和房間移開。後面的重量差要到粉體路徑上去找。", "size": 14, "color": MUTED}])
s.shapes.add_picture("/tmp/ppt-hobo.jpg", Inches(0.42), Inches(1.6), Inches(3.3), Inches(5.2))
rect(s, 3.95, 1.6, 4.4, 5.2, OK_BG)
add_text(s, 4.15, 1.8, 4.05, 4.8, [
    {"text": "環境 · 規格內", "size": 14, "bold": True, "color": OK},
    {"text": "20.67°C    31.6% RH", "size": 22, "bold": True, "after": 10},
    {"text": "溫度規格 19–27°C，讀值在中間。", "size": 14},
    {"text": "濕度規格 30–40% RH。31.6% 合格，距下限只有 1.6 個百分點，屬於偏乾的一側。", "size": 14},
    {"text": "記錄器正在 MEMORY / LOGGING。", "size": 14},
    {"text": "室濕沒有超規，不能把沾料解釋成房間太濕。", "size": 14},
])
rect(s, 8.55, 1.6, 4.35, 5.2, OK_BG)
add_text(s, 8.75, 1.8, 4.0, 4.8, [
    {"text": "ISHIDA · 測試前合格", "size": 14, "bold": True, "color": OK},
    {"text": "Unit 5，600 g 校正", "size": 22, "bold": True, "after": 10},
    {"text": "重量確認 200 g、400 g、600 g，結果合格。", "size": 14},
    {"text": "200 g 就是這一輪的名義袋重，靜態砝碼過得了。", "size": 14},
    {"text": "前三袋和 MT 對不上，不是這一顆負荷元沒校正。", "size": 14},
    {"text": "校正對的是砝碼。結塊、沾壁、落料不完整，砝碼測不到。", "size": 14},
])
footer(s, 2, "SDA 1L · 2026/10/01")

# ---------- 3 overview ----------
s = blank(); kicker(s, "總覽"); title(s, "左：問題          右：根因")
rows = [
    ("1", "前三袋，Ishida 與 MT 重量不一致",
     "Ishida 明細 203.4、200.9、201.1 g。第 1 袋比 200 g 重 3.4 g，是 50 袋裡唯一超過 +2 g 的點。第 2、第 3 袋已回到這一批的正常帶。",
     "秤斗裡的粉沒有全部進袋",
     "結塊與沾壁讓 Ishida 已扣掉的重量，和 MT 秤到的袋重差一塊殘留。第 1 袋另外帶了啟動時一次掉下的殘料。MT 三筆讀值這包資料沒有，差值還算不出來。"),
    ("2", "SDA 結塊堵在 cut gate",
     "現場把振動時間加長。影片沒有拍到閘門卡住的瞬間，拍到的是閘門上游的秤斗裡已經有硬塊。",
     "硬塊在閘口搭橋，原振動時間不夠",
     "原料裡的塊不容易過閘，細粉又黏，塊一座住就搭橋。加長振動可以把這一輪的橋震開，但不會把原料裡的塊清掉。"),
    ("3", "粉沾在 cut gate 與 weight hopper",
     "11:57 拆下的落料槽，上半仍是鏡面，中線以下是霧面膜，角落最厚。進料護網右側網緣也結了一圈粉。",
     "黏性細粉被硬塊抹在不鏽鋼上",
     "膜一旦形成，下一劑黏的是那層膜。通道上沒有看到敲擊或氣刀。室濕在規格內且偏乾，照片是乾膜，不是房間結露。"),
    ("4", "原料有結塊",
     "測試前已完成 Unit 5 / 600 g 校正，以及 200、400、600 g 確認，結果合格。結塊是在合格之後、粉一進秤斗就看見的。",
     "結塊是原料帶進來的，問題從這裡開始",
     "硬塊停在秤斗粉面上方，還沒有到 cut gate。它是問題 1 到 3 的共同起點：塊去堵閘，碎塊去抹壁，殘留造成兩台秤不一樣。"),
]
top = 1.15
for n, ih, ib, ch, cb in rows:
    rect(s, 0.42, top, 0.42, 1.35, NAVY)
    add_text(s, 0.42, top + 0.42, 0.42, 0.5, [{"text": n, "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])
    rect(s, 0.84, top, 6.0, 1.35, ISSUE_BG)
    add_text(s, 0.96, top + 0.06, 5.75, 1.22, [
        {"text": ih, "size": 13, "bold": True, "color": ISSUE, "after": 1},
        {"text": ib, "size": 11, "color": INK},
    ])
    rect(s, 6.84, top, 6.05, 1.35, CAUSE_BG)
    add_text(s, 6.96, top + 0.06, 5.8, 1.22, [
        {"text": ch, "size": 13, "bold": True, "color": CAUSE, "after": 1},
        {"text": cb, "size": 11, "color": INK},
    ])
    top += 1.45
footer(s, 3, "名義目標 200 g · 明細見第 6 頁")

# ---------- 4 chain ----------
s = blank(); kicker(s, "一條鏈"); title(s, "四個問題怎麼串在一起")
add_text(s, 0.42, 1.08, 12, 0.3, [{"text": "校正與環境停在鏈的外面。它們合格，所以沒有放進這五步。", "size": 13, "color": MUTED}])
steps = [
    ("01 原料", "秤斗裡就有硬塊", "影片裡塊在粉面上，形狀完整，不是閘門切出來的。"),
    ("02 閘門", "塊在 cut gate 搭橋", "現場紀錄。原振動時間打不開這座橋。"),
    ("03 處置", "振動時間加長", "能把橋震開。也會把黏粉再抹開，不能取代過篩。"),
    ("04 沾料", "閘門與秤斗留膜", "11:57 落料槽下半段，以及護網網緣。"),
    ("05 袋重", "Ishida ≠ MT", "第 1 袋 203.4 g。其後 49 袋收斂到 200.98 g。"),
]
for i, (a, b, c) in enumerate(steps):
    x = 0.42 + i * 2.56
    rect(s, x, 1.5, 2.42, 1.7, WHITE)
    add_text(s, x + 0.1, 1.58, 2.22, 1.55, [
        {"text": a, "size": 11, "bold": True, "color": ISSUE, "after": 2},
        {"text": b, "size": 14, "bold": True, "after": 2},
        {"text": c, "size": 11, "color": MUTED},
    ])
rect(s, 0.42, 3.38, 12.5, 0.85, NAVY)
add_text(s, 0.58, 3.48, 12.2, 0.68, [{"text": "負荷元回答的是「斗裡現在有多重」。MT 回答的是「袋裡最後有多重」。中間有粉留下，兩個數字就會分開。", "size": 14, "color": WHITE}])
s.shapes.add_picture("/tmp/ppt-lumps-wide.jpg", Inches(0.42), Inches(4.4), Inches(4.0), Inches(2.5))
s.shapes.add_picture("/tmp/ppt-grate.jpg", Inches(4.62), Inches(4.4), Inches(4.0), Inches(2.5))
s.shapes.add_picture("/tmp/ppt-chute.jpg", Inches(8.82), Inches(4.4), Inches(4.1), Inches(2.5))
footer(s, 4, "已做的處置只覆蓋第 03 步 · 影片請用網頁版播放")

# ---------- 5 issue 1 ----------
s = blank(); kicker(s, "ISSUE 1 · 重量"); title(s, "前三袋對不上，離群的是第 1 袋")
stats = [("203.4 g", "第 1 袋，比目標 +3.4 g"), ("200.9 / 201.1", "第 2、第 3 袋"), ("201.03 g", "50 袋平均，標準差 0.60"), ("200.98 g", "去掉第 1 袋，標準差 0.50"), ("49 / 50", "落在 200 ± 2 g 以內")]
for i, (a, b) in enumerate(stats):
    x = 0.42 + i * 2.56
    rect(s, x, 1.2, 2.45, 0.85, WHITE)
    add_text(s, x + 0.1, 1.26, 2.25, 0.72, [{"text": a, "size": 18, "bold": True, "after": 0}, {"text": b, "size": 11, "color": MUTED}])
rect(s, 0.42, 2.2, 6.2, 4.7, ISSUE_BG)
add_text(s, 0.58, 2.32, 5.9, 4.45, [
    {"text": "ISSUE", "size": 12, "bold": True, "color": ISSUE},
    {"text": "前三袋用 Ishida 與 MT 各秤一次，兩邊數字不一樣。", "size": 13},
    {"text": "Ishida 這一側：203.4、200.9、201.1 g。前三袋平均 201.80 g；第 4 到第 50 袋平均 200.98 g。差距幾乎全部來自第 1 袋。", "size": 13},
    {"text": "全距 4.0 g（199.4–203.4）。低於 200 g 的只有第 5 袋 199.8 g、第 6 袋 199.4 g。", "size": 13},
    {"text": "畫面 Average、S.D.、MAX、MIN、Range 都是 0，那一欄沒有計算。統計依 50 筆明細重算。", "size": 13},
    {"text": "T.Weight 11.2 kg、T.Count 56。50 袋合計 10.052 kg。多出的 6 次沒有袋號。", "size": 13},
    {"text": "CLR 10:44，機器時間 11:38，照片時間 11:27。書面公差沒有提供，這裡只描述對 200 g 的分布。", "size": 12, "color": MUTED},
])
rect(s, 6.8, 2.2, 6.1, 4.7, CAUSE_BG)
add_text(s, 6.96, 2.32, 5.8, 4.45, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "看見了：第 1 袋明顯偏重；落料槽與護網有殘留；秤斗粉面有硬塊。", "size": 13},
    {"text": "推論：Ishida 在放料前就把斗內重量算進這一袋。粉若黏在 cut gate 或斗壁，MT 就少掉那一塊；下一袋把殘料抖下來，就變成多。第 1 袋 +3.4 g 符合啟動時積住的料一次掉下。", "size": 13},
    {"text": "推論：第 2、第 3 袋若 MT 仍對不上，比較像固定偏差，空袋皮重不一致，或每一袋留下一層薄膜。", "size": 13},
    {"text": "還缺：MT 三筆讀值、空袋皮重、兩台秤是不是同一袋。沒有這三項，差值的克數不能寫死。", "size": 13},
    {"text": "排除：Unit 5 在 600 g 校正過，200 / 400 / 600 g 確認合格。負荷元不是根因。", "size": 13},
])
footer(s, 5, "明細與曲線在下一頁")

# ---------- 6 data ----------
s = blank(); kicker(s, "ISSUE 1 · 50 袋明細"); title(s, "啟動一袋偏高，之後貼著 201 g")
s.shapes.add_picture("/tmp/sda-chart.png", Inches(0.35), Inches(1.15), Inches(12.6), Inches(3.15))
add_text(s, 0.45, 4.3, 12.4, 0.3, [{"text": "橘點高於 201.5 g。藍點低於 200 g，只有第 5、第 6 袋。橘線是 200 g。相對 200 g：±1 g 有 31 袋（62%），±2 g 有 49 袋（98%）。", "size": 12, "color": MUTED}])
# mini table of 50 as text rows of 10
lines = []
for r in range(5):
    chunk = W[r * 10:(r + 1) * 10]
    lines.append("   ".join(f"{r*10+i+1:02d} {v:.1f}" for i, v in enumerate(chunk)))
add_text(s, 0.45, 4.65, 12.5, 2.2, [{"text": t, "size": 11, "after": 2} for t in lines])
footer(s, 6, "來源：Ishida 3 Log 1 window。統計欄為 0，未採用。")

# ---------- 7 issue 2 ----------
s = blank(); kicker(s, "ISSUE 2 · CUT GATE"); title(s, "結塊堵閘。振動加長是臨時處置，不是根因對策")
s.shapes.add_picture("/tmp/ppt-lumps-tall.jpg", Inches(0.4), Inches(1.25), Inches(3.15), Inches(5.6))
rect(s, 3.7, 1.25, 4.6, 5.6, ISSUE_BG)
add_text(s, 3.85, 1.4, 4.3, 5.3, [
    {"text": "ISSUE", "size": 12, "bold": True, "color": ISSUE},
    {"text": "SDA 結塊堵在 cut gate。", "size": 14},
    {"text": "現場已把振動時間加長，用來把堵住的料送出去。", "size": 14},
    {"text": "左圖是秤斗粉面上的硬塊。塊的邊緣清楚，沒有跟著細粉攤平。", "size": 14},
    {"text": "影片約 6 秒，機器是停的，沒有拍到閘門卡住的那一下。卡住是現場紀錄。", "size": 14},
    {"text": "影片檔：media/lumps.mp4，網頁版可直接播放。", "size": 12, "color": MUTED},
])
rect(s, 8.45, 1.25, 4.45, 5.6, CAUSE_BG)
add_text(s, 8.6, 1.4, 4.15, 5.3, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "看見了：硬塊在進閘之前就在秤斗裡。", "size": 13},
    {"text": "推論：細粉有黏性，填進塊和閘口的縫，搭成一座橋。後面的料下不來，閘也關不乾淨。", "size": 13},
    {"text": "推論：原來的振動時間短於把這一批的橋震開所需要的時間。", "size": 13},
    {"text": "限度：加長振動不減少原料裡的塊。時間太長，斗裡的料會被多送，黏粉會被再抹開。11:27 的 50 筆看不出是改參數之前還是之後。", "size": 13},
    {"text": "改完之後要另跑一輪，確認第 1 袋不再是 203.4 g。", "size": 13},
])
footer(s, 7, "影片 VID-20261001-WA0006")

# ---------- 8 issue 3 ----------
s = blank(); kicker(s, "ISSUE 3 · 沾料"); title(s, "粉黏在落料路徑上，留下的就是袋重差")
s.shapes.add_picture("/tmp/ppt-chute-tall.jpg", Inches(0.4), Inches(1.25), Inches(3.5), Inches(4.55))
s.shapes.add_picture("/tmp/ppt-grate-tall.jpg", Inches(4.05), Inches(1.25), Inches(3.3), Inches(4.55))
add_text(s, 0.4, 5.85, 6.9, 1.05, [
    {"text": "左：11:57 落料槽，上半鏡面，中線以下是膜。右：護網右側網緣結粉。", "size": 12, "color": MUTED},
    {"text": "護網影片 media/grate.mp4，約 10 秒，網頁版可播放。", "size": 12, "color": MUTED},
])
rect(s, 7.55, 1.25, 5.35, 5.6, CAUSE_BG)
add_text(s, 7.7, 1.4, 5.05, 5.3, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "看見了：殘留是抹上去的膜。上半還沒有、下半整片發霧，粉是沿著流動的下半段被抹開。", "size": 13},
    {"text": "推論：硬塊過閘口時被擠碎，把黏的細粉塗在金屬上。膜形成後，新的粉黏著膜。", "size": 13},
    {"text": "看見了：護網只在一側結粉。料在那一側停過。", "size": 13},
    {"text": "排除：HOBO 31.6% RH，在 30–40% 內且偏乾。槽壁沒有水珠。房間濕度不是這層膜的原因。", "size": 13},
    {"text": "還要查：原料水分，以及清潔後槽是否完全乾。進料前先拍一張空槽。", "size": 13},
    {"text": "這層膜就是 Issue 1 的機制：斗裡秤到的克數，有一部分沒有進袋。", "size": 13, "bold": True},
])
footer(s, 8, "落料槽 11:57 · 護網影片 VID-20261001-WA0004")

# ---------- 9 issue 4 ----------
s = blank(); kicker(s, "ISSUE 4 · 原料"); title(s, "結塊在秤斗裡就看見了。校正合格，所以原因在料")
s.shapes.add_picture("/tmp/ppt-lumps-wide2.jpg", Inches(0.4), Inches(1.25), Inches(3.3), Inches(5.6))
rect(s, 3.9, 1.25, 4.5, 5.6, ISSUE_BG)
add_text(s, 4.05, 1.4, 4.2, 5.3, [
    {"text": "ISSUE", "size": 12, "bold": True, "color": ISSUE},
    {"text": "原料有結塊。硬塊是立方和片狀，停在細粉錐上面，沒有陷進去。", "size": 14},
    {"text": "這是進 cut gate 之前的狀態。塊不是閘門製造出來的。", "size": 14},
    {"text": "同一天、測試開始前，Ishida 已做完兩件事且合格：Unit 5 以 600 g 校正；200、400、600 g 重量確認。", "size": 14},
])
rect(s, 8.55, 1.25, 4.35, 5.6, CAUSE_BG)
add_text(s, 8.7, 1.4, 4.05, 5.3, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "看見了：結塊進系統時就是塊。秤斗沒有把它們打碎。", "size": 13},
    {"text": "看見了：靜態砝碼在 200、400、600 g 都合格。同一台秤，砝碼過、實粉不過，差在粉的流動。", "size": 13},
    {"text": "推論：這一批原料沒有在進 Ishida 之前過篩或碎塊。振動時間是閘門堵住之後才加的，屬於下游補救。", "size": 13},
    {"text": "共同原因：問題 1、2、3 都從這裡出去。", "size": 13, "bold": True},
])
footer(s, 9, "校正為現場紀錄 · 結塊為影片 WA0006")

# ---------- 10 actions ----------
s = blank(); kicker(s, "下一步"); title(s, "已經做的，和還沒有堵住根因的")
rect(s, 0.42, 1.25, 6.15, 4.55, WHITE)
add_text(s, 0.6, 1.4, 5.8, 4.25, [
    {"text": "這一輪已經做了", "size": 16, "bold": True, "color": OK},
    {"text": "Unit 5 以 600 g 校正，200 / 400 / 600 g 確認合格。負荷元排除了。", "size": 14},
    {"text": "Cut gate 被堵住之後，振動時間已加長。這是臨時處置。", "size": 14},
    {"text": "環境有記錄：20.67°C、31.6% RH，在規格內。", "size": 14},
    {"text": "50 袋 Ishida 明細有留下。畫面統計欄是 0，以後以明細為準。", "size": 14},
])
rect(s, 6.75, 1.25, 6.15, 4.55, WHITE)
add_text(s, 6.95, 1.4, 5.8, 4.25, [
    {"text": "下一輪開始前要補上", "size": 16, "bold": True, "color": ISSUE},
    {"text": "1  原料進斗前過篩或碎塊。振動時間取代不了這一步。", "size": 13},
    {"text": "2  清潔後、進料前先拍空的落料槽和秤斗。10 袋後再看膜是否回來。", "size": 13},
    {"text": "3  同一袋同時記 Ishida 與 MT，寫下空袋皮重，至少前 10 袋。", "size": 13},
    {"text": "4  改完振動後重跑，確認沒有再出現 203.4 g 這種啟動袋。", "size": 13},
    {"text": "5  把 T.Count 56 和 50 袋對上，寫明多出的 6 次是什麼。", "size": 13},
    {"text": "6  批記錄加上結塊狀況與原料水分。", "size": 13},
])
rect(s, 0.42, 5.95, 12.48, 0.95, NAVY)
add_text(s, 0.6, 6.1, 12.15, 0.7, [{"text": "下一輪只看兩件事：第 1 袋不再離群；同一袋扣掉空袋皮重之後，Ishida 和 MT 的差，要對得上閘門上留下的那一層粉。", "size": 15, "color": WHITE}])
footer(s, 10, "臨時處置不等於根因已經關閉")

# ---------- 11 limits ----------
s = blank(); kicker(s, "證據與界限"); title(s, "這份簡報用了什麼、沒有用什麼")
rect(s, 0.42, 1.2, 6.2, 5.7, ISSUE_BG)
add_text(s, 0.58, 1.35, 5.9, 5.4, [
    {"text": "用了的證據", "size": 14, "bold": True, "color": ISSUE},
    {"text": "充填線照片，2026/10/01 09:28。", "size": 13},
    {"text": "HOBO：20.67°C、31.6% RH。規格 19–27°C、30–40%。", "size": 13},
    {"text": "Ishida 袋號 1–50，T.Weight 11.2 kg，T.Count 56，CLR 10:44，機器時間 11:38。統計欄為 0，已改以明細重算。", "size": 13},
    {"text": "落料槽內壁，11:57。", "size": 13},
    {"text": "影片 WA0006，6 秒，秤斗內結塊。", "size": 13},
    {"text": "影片 WA0004，10 秒，護網結粉。約第 3 秒有一句聽起來像「下料」，背景很吵，不當成正式紀錄。", "size": 13},
    {"text": "現場口述：前三袋兩秤不一致；結塊堵閘後加長振動；粉沾 cut gate 與秤斗；測試前校正與三點確認合格。", "size": 13},
])
rect(s, 6.8, 1.2, 6.1, 5.7, CAUSE_BG)
add_text(s, 6.96, 1.35, 5.8, 5.4, [
    {"text": "故意沒有寫死的部分", "size": 14, "bold": True, "color": CAUSE},
    {"text": "MT 的三筆克數。照片和影片都沒有，相差幾克不猜。", "size": 13},
    {"text": "振動時間改前、改後的秒數，以及 11:27 的 50 筆是改前還是改後。", "size": 13},
    {"text": "袋重書面公差。沒有規格，就不把 201.03 g 判成合格或不合格。", "size": 13},
    {"text": "T.Count 56 比 50 袋多出的 6 次是什麼。", "size": 13},
    {"text": "落料槽下半的細點是粉還是清潔水跡。需要進料前的空槽照片。", "size": 13},
    {"text": "兩支影片都是停機狀態，沒有拍到閘門堵住，也沒有拍到加長振動之後的落料。", "size": 13},
    {"text": "50 袋合計 10,051.6 g。平均 201.032 g，標準差 0.600 g，中位數 200.95 g。", "size": 13, "bold": True},
])
footer(s, 11, "2026/10/01 SDA 1 L 第一輪 · 問題與根因")

prs.save(OUT)
print("saved", OUT)
