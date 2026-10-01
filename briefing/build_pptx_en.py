#!/usr/bin/env python3
"""Rebuild the 2026-10-01 SDA filling-test deck. Numbers come from the Ishida log photo."""

from PIL import Image, ImageDraw, ImageFont
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

OUT = "/workspace/briefing/SDA-20261001-filling-test-en.pptx"
MEDIA = "/workspace/briefing/media"
FONT = "Calibri"
W = [203.4,200.9,201.1,200.3,199.8,199.4,201.0,201.5,200.8,201.1,201.8,200.4,201.2,200.8,200.7,200.7,201.0,201.0,200.9,201.4,200.6,200.9,201.0,201.1,201.1,200.6,201.6,201.4,201.6,202.0,201.6,202.0,201.8,201.6,201.4,201.0,200.9,200.7,201.2,200.8,201.0,200.8,200.8,200.6,200.8,200.6,200.8,200.7,200.6,200.8]

NAVY = RGBColor(0x1B, 0x2A, 0x3D)
INK = RGBColor(0x1C, 0x19, 0x15)
MUTED = RGBColor(0x5E, 0x58, 0x4F)
ISSUE = RGBColor(0x9A, 0x3B, 0x12)
CAUSE = RGBColor(0x1E, 0x4D, 0x78)
OK = RGBColor(0x0E, 0x6B, 0x45)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
PAPER = RGBColor(0xF7, 0xF4, 0xEE)
ISSUE_BG = RGBColor(0xFF, 0xF1, 0xE8)
CAUSE_BG = RGBColor(0xEE, 0xF4, 0xF8)
OK_BG = RGBColor(0xE7, 0xF5, 0xEE)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "2026/10/01 SDA 1 L filling test — issues and root causes"
prs.core_properties.subject = "First run, 50 bags"


def set_run(run, text, size, bold=False, color=INK, font=FONT):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", font)


def add_text(slide, l, t, w, h, paragraphs, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    tf = box.text_frame
    tf.auto_size = None
    tf.vertical_anchor = anchor
    for i, p in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.space_after = Pt(p.get("after", 4))
        para.space_before = Pt(p.get("before", 0))
        para.level = p.get("level", 0)
        set_run(para.add_run(), p["text"], p.get("size", 14), p.get("bold", False), p.get("color", INK))
    return box


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    # no shadow
    spPr = sh._element.spPr
    effect = spPr.find(qn("a:effectLst"))
    if effect is not None:
        spPr.remove(effect)
    return sh


def blank():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 13.333, 7.5, PAPER)
    return s


def crop_fill(path, box_w, box_h, dest):
    im = Image.open(path).convert("RGB")
    target = box_w / box_h
    src = im.width / im.height
    if src > target:
        nw = int(im.height * target)
        x = (im.width - nw) // 2
        im = im.crop((x, 0, x + nw, im.height))
    else:
        nh = int(im.width / target)
        y = max(0, (im.height - nh) // 2)
        im = im.crop((0, y, im.width, y + nh))
    im.save(dest, quality=88)
    return dest


def footer(s, n, label):
    add_text(s, 0.4, 7.12, 9, 0.28, [{"text": label, "size": 10, "color": MUTED}])
    add_text(s, 11.4, 7.12, 1.5, 0.28, [{"text": f"{n:02d}", "size": 10, "color": MUTED, "align": PP_ALIGN.RIGHT}])


def kicker(s, text):
    add_text(s, 0.42, 0.22, 10, 0.28, [{"text": text, "size": 11, "bold": True, "color": MUTED}])


def title(s, text, top=0.46):
    add_text(s, 0.42, top, 12.4, 0.55, [{"text": text, "size": 26, "bold": True}])


# ---------- chart ----------
def make_chart(path):
    font = ImageFont.truetype("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 18)
    font_s = ImageFont.truetype("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc", 15)
    im = Image.new("RGB", (1800, 520), (247, 244, 238))
    d = ImageDraw.Draw(im)
    L, T, R, B = 70, 30, 20, 40
    pw, ph = 1800 - L - R, 520 - T - B
    ymin, ymax = 198.6, 204.0

    def X(i):
        return L + i * pw / 49

    def Y(v):
        return T + (ymax - v) / (ymax - ymin) * ph

    d.line([(L, Y(200)), (1800 - R, Y(200))], fill=(196, 85, 42), width=2)
    d.line([(L, Y(201)), (1800 - R, Y(201))], fill=(217, 208, 195), width=1)
    for g in (199, 200, 201, 202, 203, 204):
        d.text((4, Y(g) - 10), str(g), font=font_s, fill=(138, 129, 117))
    pts = [(X(i), Y(v)) for i, v in enumerate(W)]
    d.line(pts, fill=(30, 77, 120), width=3)
    for i, v in enumerate(W):
        r = 6 if v > 201.5 or v < 200 else 4
        col = (196, 85, 42) if v > 201.5 else (30, 77, 120) if v < 200 else (27, 42, 61)
        d.ellipse((X(i) - r, Y(v) - r, X(i) + r, Y(v) + r), fill=col)
    d.text((X(0) + 8, Y(203.4) - 26), "Bag 1  203.4", font=font, fill=(154, 59, 18))
    for t in range(0, 50, 5):
        d.text((X(t) - 8, 490), str(t + 1), font=font_s, fill=(138, 129, 117))
    im.save(path, quality=90)


make_chart("/tmp/sda-chart-en.png")
crop_fill(f"{MEDIA}/line.jpg", 6.183, 7.5, "/tmp/ppt-line.jpg")
crop_fill(f"{MEDIA}/hobo.jpg", 3.3, 5.2, "/tmp/ppt-hobo.jpg")
crop_fill(f"{MEDIA}/lumps.jpg", 4.0, 2.5, "/tmp/ppt-lumps-wide.jpg")
crop_fill(f"{MEDIA}/grate-close.jpg", 4.0, 2.5, "/tmp/ppt-grate.jpg")
crop_fill(f"{MEDIA}/chute.jpg", 4.1, 2.5, "/tmp/ppt-chute.jpg")
crop_fill(f"{MEDIA}/lumps.jpg", 3.15, 5.6, "/tmp/ppt-lumps-tall.jpg")
crop_fill(f"{MEDIA}/chute.jpg", 3.5, 4.55, "/tmp/ppt-chute-tall.jpg")
crop_fill(f"{MEDIA}/grate-close.jpg", 3.3, 4.55, "/tmp/ppt-grate-tall.jpg")
crop_fill(f"{MEDIA}/lumps-wide.jpg", 3.3, 5.6, "/tmp/ppt-lumps-wide2.jpg")

# ---------- 1 cover ----------
s = prs.slides.add_slide(prs.slide_layouts[6])
rect(s, 0, 0, 7.15, 7.5, NAVY)
rect(s, 7.15, 0, 6.183, 7.5, NAVY)
s.shapes.add_picture("/tmp/ppt-line.jpg", Inches(7.15), Inches(0), Inches(6.183), Inches(7.5))
add_text(s, 0.5, 0.7, 6.3, 0.3, [{"text": "FILLING TEST  ·  ISSUE / ROOT CAUSE", "size": 12, "bold": True, "color": RGBColor(0xD7, 0xB4, 0x8A)}])
add_text(s, 0.5, 1.5, 6.3, 2.1, [
    {"text": "SDA 1 L filling test", "size": 36, "bold": True, "color": WHITE, "after": 0},
    {"text": "First run, 50 bags", "size": 36, "bold": True, "color": WHITE},
])
add_text(s, 0.5, 3.8, 6.2, 1.6, [{"text": "1 October 2026. All four findings are one event: lumpy powder did not leave the weigh hopper cleanly, so the gate blocked, the walls filmed over, and the weight Ishida counted did not all reach the bag.", "size": 16, "color": RGBColor(0xE6, 0xDF, 0xD2)}])
add_text(s, 0.5, 6.15, 2, 0.7, [{"text": "Nominal weight", "size": 12, "color": RGBColor(0xC9, 0xC1, 0xB4), "after": 0}, {"text": "200 g / bag", "size": 16, "bold": True, "color": WHITE}])
add_text(s, 2.6, 6.15, 2, 0.7, [{"text": "On the line", "size": 12, "color": RGBColor(0xC9, 0xC1, 0xB4), "after": 0}, {"text": "09:28 – 11:57", "size": 16, "bold": True, "color": WHITE}])
add_text(s, 4.7, 6.15, 2.2, 0.7, [{"text": "Ruled out", "size": 12, "color": RGBColor(0xC9, 0xC1, 0xB4), "after": 0}, {"text": "Load cell, room", "size": 16, "bold": True, "color": WHITE}])

# ---------- 2 ruled out ----------
s = blank(); kicker(s, "RULED OUT FIRST"); title(s, "Before the run, the scale and the room both passed")
add_text(s, 0.42, 1.1, 12, 0.35, [{"text": "Those two results move the root cause off the instrument and off the room. The weight gap is on the powder path.", "size": 14, "color": MUTED}])
s.shapes.add_picture("/tmp/ppt-hobo.jpg", Inches(0.42), Inches(1.6), Inches(3.3), Inches(5.2))
rect(s, 3.95, 1.6, 4.4, 5.2, OK_BG)
add_text(s, 4.15, 1.8, 4.05, 4.8, [
    {"text": "ROOM · INSIDE SPEC", "size": 14, "bold": True, "color": OK},
    {"text": "20.67°C    31.6% RH", "size": 22, "bold": True, "after": 10},
    {"text": "Temperature spec 19–27°C. The reading sits in the middle.", "size": 14},
    {"text": "Humidity spec 30–40% RH. 31.6% passes, 1.6 points above the floor, on the dry side.", "size": 14},
    {"text": "The logger shows MEMORY / LOGGING.", "size": 14},
    {"text": "Room humidity did not exceed the limit. The film is not a wet room.", "size": 14},
])
rect(s, 8.55, 1.6, 4.35, 5.2, OK_BG)
add_text(s, 8.75, 1.8, 4.0, 4.8, [
    {"text": "ISHIDA · PASSED BEFORE THE RUN", "size": 14, "bold": True, "color": OK},
    {"text": "Unit 5, 600 g calibration", "size": 22, "bold": True, "after": 10},
    {"text": "Weight checks at 200 g, 400 g, and 600 g all passed.", "size": 14},
    {"text": "200 g is this run’s nominal bag weight. A static weight passed it.", "size": 14},
    {"text": "The first three bags missing MT is not an uncalibrated load cell.", "size": 14},
    {"text": "Calibration checks a weight. Lumps, film, and a short dump are invisible to it.", "size": 14},
])
footer(s, 2, "SDA 1 L · 1 Oct 2026")

# ---------- 3 overview ----------
s = blank(); kicker(s, "OVERVIEW"); title(s, "Left: issue.          Right: root cause.")
rows = [
    ("1", "First three bags: Ishida and MT do not match",
     "Ishida reads 203.4, 200.9, 201.1 g. Bag 1 is +3.4 g versus 200 g, the only bag beyond +2 g. Bags 2 and 3 are back in the pack.",
     "Hopper powder did not all reach the bag",
     "Lumps and wall film separate the weight Ishida counted from the weight MT sees. Bag 1 also carried a startup slug. The three MT readings are not in this pack, so the gram gap is not calculated."),
    ("2", "SDA lumps blocked the cut gate",
     "Vibration time was lengthened. The video does not show the jam. It shows hard lumps already in the hopper, upstream of the gate.",
     "Lumps bridged the opening. Vibration was too short.",
     "Lumps from the raw material seat in the gate, and sticky fines build a bridge. Longer vibration breaks this bridge. It does not remove the lumps."),
    ("3", "Powder stuck on the cut gate and weigh hopper",
     "At 11:57 the removed chute is bright above the midline and filmed below it, thickest in the corners. The infeed grate is caked on the right rim only.",
     "Sticky fines were smeared onto stainless by the lumps",
     "Once a film exists, the next dose sticks to the film. No knocker or air knife is visible. Room RH is inside spec and on the dry side. This is a dry film, not condensation."),
    ("4", "The raw material contains lumps",
     "Before the test, Unit 5 was calibrated at 600 g and the 200 / 400 / 600 g checks passed. The lumps were seen after that pass, as soon as powder entered the hopper.",
     "The lumps came in with the raw material",
     "They sit on the powder surface, before the cut gate. They start issues 1 to 3: lumps block the gate, broken lumps smear the wall, residue splits the two scales."),
]
top = 1.15
for n, ih, ib, ch, cb in rows:
    rect(s, 0.42, top, 0.42, 1.35, NAVY)
    add_text(s, 0.42, top + 0.42, 0.42, 0.5, [{"text": n, "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}])
    rect(s, 0.84, top, 6.0, 1.35, ISSUE_BG)
    add_text(s, 0.96, top + 0.06, 5.75, 1.22, [
        {"text": ih, "size": 13, "bold": True, "color": ISSUE, "after": 1},
        {"text": ib, "size": 11, "color": INK},
    ])
    rect(s, 6.84, top, 6.05, 1.35, CAUSE_BG)
    add_text(s, 6.96, top + 0.06, 5.8, 1.22, [
        {"text": ch, "size": 13, "bold": True, "color": CAUSE, "after": 1},
        {"text": cb, "size": 11, "color": INK},
    ])
    top += 1.45
footer(s, 3, "Nominal target 200 g · detail on page 6")

# ---------- 4 chain ----------
s = blank(); kicker(s, "ONE CHAIN"); title(s, "How the four issues connect")
add_text(s, 0.42, 1.08, 12, 0.3, [{"text": "Calibration and the room sit outside this chain. They passed, so they are not one of the five steps.", "size": 13, "color": MUTED}])
steps = [
    ("01  RAW MATERIAL", "Lumps are already in the hopper", "Whole pieces sit on the powder. The gate did not cut them."),
    ("02  GATE", "Lumps bridge the cut gate", "Line record. The original vibration time does not break that bridge."),
    ("03  ACTION", "Vibration time increased", "It breaks the bridge and smears more powder. It does not replace a screen."),
    ("04  FILM", "Gate and hopper keep a film", "Lower chute at 11:57, and the grate rim. That is weight that missed the bag."),
    ("05  BAG WEIGHT", "Ishida ≠ MT", "Bag 1 is 203.4 g. The next 49 bags settle at 200.98 g."),
]
for i, (a, b, c) in enumerate(steps):
    x = 0.42 + i * 2.56
    rect(s, x, 1.5, 2.42, 1.7, WHITE)
    add_text(s, x + 0.1, 1.58, 2.22, 1.55, [
        {"text": a, "size": 11, "bold": True, "color": ISSUE, "after": 2},
        {"text": b, "size": 14, "bold": True, "after": 2},
        {"text": c, "size": 11, "color": MUTED},
    ])
rect(s, 0.42, 3.38, 12.5, 0.85, NAVY)
add_text(s, 0.58, 3.48, 12.2, 0.68, [{"text": "The load cell answers how much is in the hopper. MT answers how much ended up in the bag. Powder left on the gate or the wall splits those two numbers.", "size": 15, "color": WHITE}])
s.shapes.add_picture("/tmp/ppt-lumps-wide.jpg", Inches(0.42), Inches(4.4), Inches(4.0), Inches(2.5))
s.shapes.add_picture("/tmp/ppt-grate.jpg", Inches(4.62), Inches(4.4), Inches(4.0), Inches(2.5))
s.shapes.add_picture("/tmp/ppt-chute.jpg", Inches(8.82), Inches(4.4), Inches(4.1), Inches(2.5))
footer(s, 4, "The action taken covers step 03 only · play the videos in the HTML deck")

# ---------- 5 issue 1 ----------
s = blank(); kicker(s, "ISSUE 1 · WEIGHT"); title(s, "The first three bags disagree. Bag 1 is the outlier.")
stats = [("203.4 g", "Bag 1, +3.4 g vs target"), ("200.9 / 201.1", "Bags 2 and 3"), ("201.03 g", "Mean of 50, SD 0.60"), ("200.98 g", "Without bag 1, SD 0.50"), ("49 / 50", "Inside 200 ± 2 g")]
for i, (a, b) in enumerate(stats):
    x = 0.42 + i * 2.56
    rect(s, x, 1.2, 2.45, 0.85, WHITE)
    add_text(s, x + 0.1, 1.26, 2.25, 0.72, [{"text": a, "size": 18, "bold": True, "after": 0}, {"text": b, "size": 11, "color": MUTED}])
rect(s, 0.42, 2.2, 6.2, 4.7, ISSUE_BG)
add_text(s, 0.58, 2.32, 5.9, 4.45, [
    {"text": "ISSUE", "size": 12, "bold": True, "color": ISSUE},
    {"text": "The first three bags were weighed on Ishida and on MT. The two numbers do not match.", "size": 13},
    {"text": "Ishida side: 203.4, 200.9, 201.1 g. Mean of the first three is 201.80 g. Mean of bags 4–50 is 200.98 g. Almost all of that gap is bag 1.", "size": 13},
    {"text": "Range is 4.0 g (199.4–203.4). Only bag 5 (199.8 g) and bag 6 (199.4 g) are under 200 g.", "size": 13},
    {"text": "On-screen Average, S.D., MAX, MIN, and Range are 0. That column was not calculated. Figures below are from the 50 lines.", "size": 13},
    {"text": "T.Weight 11.2 kg, T.Count 56. The 50 bags add up to 10.052 kg. The extra 6 dumps have no bag number.", "size": 13},
    {"text": "CLR 10:44, machine time 11:38, photo time 11:27. No written tolerance was provided. This is the spread around 200 g.", "size": 12, "color": MUTED},
])
rect(s, 6.8, 2.2, 6.1, 4.7, CAUSE_BG)
add_text(s, 6.96, 2.32, 5.8, 4.45, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "Seen: bag 1 is heavy. The chute and grate hold residue. Hard lumps sit on the hopper powder.", "size": 13},
    {"text": "Inferred: Ishida counts the hopper before discharge. Powder stuck on the gate or wall is missing from the bag MT weighs. The next bag can run heavy if that residue lets go. Bag 1 at +3.4 g fits a startup slug.", "size": 13},
    {"text": "Inferred: if bags 2 and 3 still miss MT, that looks like a fixed offset — tare, or a thin film on every bag.", "size": 13},
    {"text": "Missing: the three MT readings, the empty-bag tare, and proof both scales weighed the same bag. The gram gap cannot be stated.", "size": 13},
    {"text": "Ruled out: Unit 5 was calibrated at 600 g. The 200 / 400 / 600 g checks passed. The load cell is not the cause.", "size": 13},
])
footer(s, 5, "Detail and the curve are on the next page")

# ---------- 6 data ----------
s = blank(); kicker(s, "ISSUE 1 · 50-BAG DETAIL"); title(s, "One startup bag runs high. The rest sit near 201 g.")
s.shapes.add_picture("/tmp/sda-chart-en.png", Inches(0.35), Inches(1.15), Inches(12.6), Inches(3.15))
add_text(s, 0.45, 4.3, 12.4, 0.3, [{"text": "Orange points are above 201.5 g. Blue points are below 200 g: bags 5 and 6 only. The orange line is 200 g. Versus 200 g: ±1 g holds 31 bags (62%), ±2 g holds 49 (98%).", "size": 12, "color": MUTED}])
# mini table of 50 as text rows of 10
lines = []
for r in range(5):
    chunk = W[r * 10:(r + 1) * 10]
    lines.append("   ".join(f"{r*10+i+1:02d} {v:.1f}" for i, v in enumerate(chunk)))
add_text(s, 0.45, 4.65, 12.5, 2.2, [{"text": t, "size": 11, "after": 2} for t in lines])
footer(s, 6, "Source: Ishida 3 Log 1 window. The statistics column is 0 and was not used.")

# ---------- 7 issue 2 ----------
s = blank(); kicker(s, "ISSUE 2 · CUT GATE"); title(s, "Lumps blocked the gate. Longer vibration is containment, not the fix.")
s.shapes.add_picture("/tmp/ppt-lumps-tall.jpg", Inches(0.4), Inches(1.25), Inches(3.15), Inches(5.6))
rect(s, 3.7, 1.25, 4.6, 5.6, ISSUE_BG)
add_text(s, 3.85, 1.4, 4.3, 5.3, [
    {"text": "ISSUE", "size": 12, "bold": True, "color": ISSUE},
    {"text": "SDA lumps blocked the cut gate.", "size": 14},
    {"text": "Vibration time was lengthened to move the blocked powder.", "size": 14},
    {"text": "The photo shows hard lumps on the hopper powder. The edges stay sharp.", "size": 14},
    {"text": "The clip is about 6 seconds and the machine is stopped. The jam itself is a line record.", "size": 14},
    {"text": "Play media/lumps.mp4 in the HTML deck.", "size": 12, "color": MUTED},
])
rect(s, 8.45, 1.25, 4.45, 5.6, CAUSE_BG)
add_text(s, 8.6, 1.4, 4.15, 5.3, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "Seen: the lumps are in the hopper before the gate.", "size": 13},
    {"text": "Inferred: sticky fines pack the gap around a lump and build a bridge. Flow stops, and the gate does not close clean.", "size": 13},
    {"text": "Inferred: the original vibration time was shorter than this lot needed to break the bridge.", "size": 13},
    {"text": "Limit: longer vibration does not remove lumps from the raw material. Too long, the hopper is overfed and powder is smeared further. The 11:27 lines do not show before versus after.", "size": 13},
    {"text": "Rerun after the change. Bag 1 should not be 203.4 g again.", "size": 13},
])
footer(s, 7, "Video VID-20261001-WA0006")

# ---------- 8 issue 3 ----------
s = blank(); kicker(s, "ISSUE 3 · STICKING"); title(s, "Powder stuck on the discharge path. That film is the bag-weight gap.")
s.shapes.add_picture("/tmp/ppt-chute-tall.jpg", Inches(0.4), Inches(1.25), Inches(3.5), Inches(4.55))
s.shapes.add_picture("/tmp/ppt-grate-tall.jpg", Inches(4.05), Inches(1.25), Inches(3.3), Inches(4.55))
add_text(s, 0.4, 5.85, 6.9, 1.05, [
    {"text": "Left: chute at 11:57, bright above the midline, filmed below. Right: grate caked on the right rim.", "size": 12, "color": MUTED},
    {"text": "Play media/grate.mp4 in the HTML deck, about 10 seconds.", "size": 12, "color": MUTED},
])
rect(s, 7.55, 1.25, 5.35, 5.6, CAUSE_BG)
add_text(s, 7.7, 1.4, 5.05, 5.3, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "Seen: the residue is a smeared film. The upper half is bare and the lower half is fogged, along the flow path.", "size": 13},
    {"text": "Inferred: lumps are crushed at the gate and paint sticky fines onto the metal. New powder then sticks to the film.", "size": 13},
    {"text": "Seen: the grate is caked on one side only. Powder stopped there.", "size": 13},
    {"text": "Ruled out: HOBO 31.6% RH is inside 30–40% and on the dry side. No droplets. Room humidity is not this film.", "size": 13},
    {"text": "Still to check: raw-material moisture, and whether the chute was dry after cleaning. Photograph it empty before the next feed.", "size": 13},
    {"text": "This film is the mechanism of issue 1: grams counted in the hopper never entered the bag.", "size": 13, "bold": True},
])
footer(s, 8, "Chute at 11:57 · grate video VID-20261001-WA0004")

# ---------- 9 issue 4 ----------
s = blank(); kicker(s, "ISSUE 4 · RAW MATERIAL"); title(s, "Lumps are already in the hopper. Calibration passed, so the cause is the material.")
s.shapes.add_picture("/tmp/ppt-lumps-wide2.jpg", Inches(0.4), Inches(1.25), Inches(3.3), Inches(5.6))
rect(s, 3.9, 1.25, 4.5, 5.6, ISSUE_BG)
add_text(s, 4.05, 1.4, 4.2, 5.3, [
    {"text": "ISSUE", "size": 12, "bold": True, "color": ISSUE},
    {"text": "The raw material contains lumps. Cubes and flakes sit on the cone of fines. They do not sink in.", "size": 14},
    {"text": "This is before the cut gate. The gate did not make these pieces.", "size": 14},
    {"text": "The same day, before the test, Ishida passed two checks: Unit 5 at 600 g, and 200 / 400 / 600 g.", "size": 14},
])
rect(s, 8.55, 1.25, 4.35, 5.6, CAUSE_BG)
add_text(s, 8.7, 1.4, 4.05, 5.3, [
    {"text": "ROOT CAUSE", "size": 12, "bold": True, "color": CAUSE},
    {"text": "Seen: the lumps entered as lumps. The hopper did not break them.", "size": 13},
    {"text": "Seen: static weights passed at 200, 400, and 600 g. Weights pass and real powder does not. The difference is flow, not the load cell.", "size": 13},
    {"text": "Inferred: this lot was not screened or delumped before the Ishida. Vibration time was added after the gate was already blocked.", "size": 13},
    {"text": "Common cause: issues 1, 2, and 3 leave from here.", "size": 13, "bold": True},
])
footer(s, 9, "Calibration is the line record · lumps are video WA0006")

# ---------- 10 actions ----------
s = blank(); kicker(s, "NEXT STEP"); title(s, "What is already done, and what has not closed the cause")
rect(s, 0.42, 1.25, 6.15, 4.55, WHITE)
add_text(s, 0.6, 1.4, 5.8, 4.25, [
    {"text": "Already done this run", "size": 16, "bold": True, "color": OK},
    {"text": "Unit 5 calibrated at 600 g. Checks at 200 / 400 / 600 g passed. The load cell is cleared.", "size": 14},
    {"text": "After the cut gate blocked, vibration time was lengthened. That is containment.", "size": 14},
    {"text": "The room was logged: 20.67°C, 31.6% RH, inside spec.", "size": 14},
    {"text": "The 50 Ishida lines were kept. On-screen statistics are 0. Use the lines.", "size": 14},
])
rect(s, 6.75, 1.25, 6.15, 4.55, WHITE)
add_text(s, 6.95, 1.4, 5.8, 4.25, [
    {"text": "Close these before the next run", "size": 16, "bold": True, "color": ISSUE},
    {"text": "1  Screen or delump before the hopper. Vibration time does not replace that.", "size": 13},
    {"text": "2  Photograph the empty chute and hopper after cleaning. Check again after 10 bags.", "size": 13},
    {"text": "3  Record Ishida and MT on the same bag, with empty-bag tare, for at least 10 bags.", "size": 13},
    {"text": "4  Rerun after the vibration change. Bag 1 should not be another 203.4 g.", "size": 13},
    {"text": "5  Reconcile T.Count 56 with the 50 bags. Name the extra 6 dumps.", "size": 13},
    {"text": "6  Add lump condition and raw-material moisture to the batch record.", "size": 13},
])
rect(s, 0.42, 5.95, 12.48, 0.95, NAVY)
add_text(s, 0.6, 6.1, 12.15, 0.7, [{"text": "Next run, two checks only: bag 1 is no longer an outlier, and on the same bag, after tare, the Ishida–MT gap matches the powder left on the gate.", "size": 15, "color": WHITE}])
footer(s, 10, "Containment is not the same as a closed root cause")

# ---------- 11 limits ----------
s = blank(); kicker(s, "EVIDENCE AND LIMITS"); title(s, "What this briefing uses, and what it leaves open")
rect(s, 0.42, 1.2, 6.2, 5.7, ISSUE_BG)
add_text(s, 0.58, 1.35, 5.9, 5.4, [
    {"text": "Evidence used", "size": 14, "bold": True, "color": ISSUE},
    {"text": "Filling-line photo, 1 Oct 2026, 09:28.", "size": 13},
    {"text": "HOBO: 20.67°C, 31.6% RH. Limits 19–27°C and 30–40%.", "size": 13},
    {"text": "Ishida bags 1–50, T.Weight 11.2 kg, T.Count 56, CLR 10:44, machine time 11:38. Statistics are 0; figures were recomputed from the lines.", "size": 13},
    {"text": "Inside wall of the discharge chute, 11:57.", "size": 13},
    {"text": "Video WA0006, 6 seconds, lumps in the weigh hopper.", "size": 13},
    {"text": "Video WA0004, 10 seconds, powder on the infeed grate. A muffled word around 3 seconds sounds like “discharge” and is not a formal record.", "size": 13},
    {"text": "Line account: first three bags disagree; lumps blocked the gate and vibration was lengthened; powder stuck on the gate and hopper; calibration and the three-point check passed.", "size": 13},
])
rect(s, 6.8, 1.2, 6.1, 5.7, CAUSE_BG)
add_text(s, 6.96, 1.35, 5.8, 5.4, [
    {"text": "Left unsolved on purpose", "size": 14, "bold": True, "color": CAUSE},
    {"text": "The three MT gram values. They are not in the photos or videos. The gap is not guessed.", "size": 13},
    {"text": "Vibration time in seconds, before and after, and whether the 11:27 lines are before or after the change.", "size": 13},
    {"text": "A written bag-weight tolerance. Without a spec, 201.03 g is not called pass or fail.", "size": 13},
    {"text": "What the 6 extra dumps inside T.Count 56 are.", "size": 13},
    {"text": "Whether the fine specks on the lower chute are powder or dry wash marks. That needs an empty-chute photo.", "size": 13},
    {"text": "Both videos are of a stopped machine. Neither shows the jam, nor discharge after the vibration change.", "size": 13},
    {"text": "50 bags total 10,051.6 g. Mean 201.032 g, SD 0.600 g, median 200.95 g.", "size": 13, "bold": True},
])
footer(s, 11, "1 Oct 2026 · SDA 1 L first run · issues and root causes")

prs.save(OUT)
print("saved", OUT)
